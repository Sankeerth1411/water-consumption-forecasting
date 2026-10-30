import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    os.makedirs('reports', exist_ok=True)
    doc = Document()

    # Page Margins (1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)

    # Helper function for headings
    def add_chapter_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 51, 102) # Academic navy
        return p

    def add_section_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(51, 51, 51)
        return p

    def add_body_paragraph(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.5
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        return p

    # --- COVER PAGE ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(12)
    r = title_p.add_run("NITTE MEENAKSHI INSTITUTE OF TECHNOLOGY\n(An Autonomous Institution, Affiliated to VTU, Belagavi)\nDepartment of Artificial Intelligence and Data Science\nYelahanka, Bengaluru - 560064")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True

    doc.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.CENTER

    proj_title_p = doc.add_paragraph()
    proj_title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    proj_title_p.paragraph_format.space_before = Pt(24)
    proj_title_p.paragraph_format.space_after = Pt(24)
    r2 = proj_title_p.add_run("A DATA MINING INDIVIDUAL PRACTICAL CASE STUDY REPORT ON\n\n\"HOUSEHOLD WATER CONSUMPTION FORECASTING USING DECISION TREE REGRESSION\"")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(18)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0, 51, 102)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_before = Pt(12)
    sub_p.paragraph_format.space_after = Pt(36)
    r3 = sub_p.add_run("Submitted in partial fulfillment of the requirements for the course\nDATA MINING (Course Code: DATA MINING-22ADG53)\n\nAcademic Year: 2026–2027")
    r3.font.name = 'Times New Roman'
    r3.font.size = Pt(12)

    # Details table
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    details = [
        ("Submitted By:", "R. Sankeerth (USN: 1NT24AD047)"),
        ("Semester / Section:", "5th Semester / Artificial Intelligence & Data Science"),
        ("Assigned Algorithm:", "Decision Tree Regressor (sklearn.tree.DecisionTreeRegressor)"),
        ("Faculty Guide:", "Mrs Nethravathi T L, Assistant Professor G-1"),
        ("Department:", "Department of Artificial Intelligence and Data Science")
    ]
    for i, (k, v) in enumerate(details):
        row = table.rows[i]
        row.cells[0].text = k
        row.cells[1].text = v
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        for cell in row.cells:
            for p in cell.paragraphs:
                p.runs[0].font.name = 'Times New Roman'
                p.runs[0].font.size = Pt(11)

    doc.add_page_break()

    # --- ABSTRACT ---
    add_chapter_heading("Abstract")
    add_body_paragraph(
        "Accurate household water consumption forecasting is vital for municipal water resource management, demand response planning, "
        "and sustainable infrastructure operations. This case study formulates daily water consumption estimation as a supervised regression "
        "task using historical residential monitoring data. Following rigorous data cleaning, temporal aggregation, and feature engineering "
        "(incorporating lagged consumption and rolling statistics), a Decision Tree Regressor (DecisionTreeRegressor) was trained and evaluated "
        "against naive persistence baselines and linear regression. The model achieved a Test Mean Absolute Error (MAE) of 88.60 Liters and "
        "Root Mean Squared Error (RMSE) of 121.71 Liters, outperforming historical persistence baselines and providing interpretable tree rules. "
        "A lightweight Streamlit interactive web application was developed to demonstrate real-time forecasting and leaf-node diagnostics."
    )

    doc.add_page_break()

    # --- TABLE OF CONTENTS ---
    add_chapter_heading("Table of Contents")
    toc_items = [
        "1. Introduction .......................................................................................................... 3",
        "2. Dataset Description ................................................................................................ 5",
        "3. Methodology and Algorithm ................................................................................. 7",
        "4. Implementation ................................................................................................... 11",
        "5. Results and Discussion ........................................................................................ 14",
        "6. Conclusion and Future Enhancements ............................................................... 18",
        "References ............................................................................................................. 20"
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)

    doc.add_page_break()

    # --- CHAPTER 1 ---
    add_chapter_heading("Chapter 1: Introduction")
    add_section_heading("1.1 Background and Real-World Importance")
    add_body_paragraph(
        "Urban water utilities face increasing stress due to rapid population growth, climate variability, and constrained water reserves. "
        "Predicting household water demand allows municipal planners and utility operators to optimize pumping schedules, detect pipe leaks early, "
        "and manage reservoir storage efficiently. Given historical water usage records across multiple residential households, the core problem "
        "addressed in this case study is estimating the total water consumption (in liters) for a household on calendar day t, using only information "
        "available prior to day t."
    )
    add_section_heading("1.2 Problem Statement")
    add_body_paragraph(
        "Given historical daily water consumption time-series data and occupancy indicators for residential households, build a predictive model "
        "using Decision Tree Regression to forecast next-day water usage accurately, without data leakage or reliance on future unobserved variables."
    )
    add_section_heading("1.3 Objectives")
    add_body_paragraph(
        "1. Audit and preprocess residential water consumption and occupancy data from Zenodo (Record 15180141).\n"
        "2. Engineer chronological lag features and rolling statistics without data leakage.\n"
        "3. Train and hyperparameter-tune a Decision Tree Regressor (sklearn.tree.DecisionTreeRegressor).\n"
        "4. Evaluate performance against established baselines (Lag-1, Lag-7, Linear Regression) using MAE, RMSE, and R².\n"
        "5. Build an interactive Streamlit demonstration application and comprehensive documentation."
    )
    add_section_heading("1.4 Scope and Applications")
    add_body_paragraph(
        "The scope encompasses 10 residential households monitored across the year 2022. The resulting models and web application can be deployed "
        "in smart-metering dashboards, domestic water conservation apps, and municipal demand forecasting systems."
    )

    # --- CHAPTER 2 ---
    add_chapter_heading("Chapter 2: Dataset Description")
    add_section_heading("2.1 Dataset Source and Citation")
    add_body_paragraph(
        "The project utilizes the publicly available 'Residential Household Dataset: Occupancy, Water, and Electricity Data' hosted on Zenodo "
        "(Record 15180141, DOI: 10.5281/zenodo.15180140). The dataset provides high-resolution monitoring data for residential buildings."
    )
    add_section_heading("2.2 Data Dictionary")
    add_body_paragraph(
        "- building: Unique building/household identifier (e.g., BUILD-D-01 to BUILD-D-10).\n"
        "- date_time: ISO-8601 timestamp (hourly interval).\n"
        "- date: Calendar date (YYYY-MM-DD).\n"
        "- water_use_l: Water consumption in liters (target/feature).\n"
        "- occ_01 to occ_05: Household occupancy indicator features.\n"
        "- lag_1: Water consumption 1 day prior (t-1).\n"
        "- lag_7: Water consumption 7 days prior (t-7).\n"
        "- rolling_mean_7: Mean water consumption over the past 7 days.\n"
        "- rolling_std_7: Standard deviation of water consumption over the past 7 days."
    )
    add_section_heading("2.3 Sample Records Table")

    # Table of sample records
    table_sample = doc.add_table(rows=4, cols=5)
    table_sample.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Building", "Date", "Water Use (L)", "Lag-1 (L)", "Rolling Mean (L)"]
    for j, h in enumerate(headers):
        table_sample.rows[0].cells[j].text = h
        table_sample.rows[0].cells[j].paragraphs[0].runs[0].font.bold = True

    sample_data = [
        ("BUILD-D-01", "2022-01-08", "275.0", "312.4", "295.1"),
        ("BUILD-D-01", "2022-01-09", "296.0", "275.0", "290.3"),
        ("BUILD-D-01", "2022-01-10", "240.0", "296.0", "288.7")
    ]
    for i, row_vals in enumerate(sample_data):
        for j, val in enumerate(row_vals):
            table_sample.rows[i+1].cells[j].text = val

    for row in table_sample.rows:
        for cell in row.cells:
            set_cell_margins(cell, 80, 80, 100, 100)
            for p in cell.paragraphs:
                if p.runs:
                    p.runs[0].font.name = 'Times New Roman'
                    p.runs[0].font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- CHAPTER 3 ---
    add_chapter_heading("Chapter 3: Methodology and Algorithm")
    add_section_heading("3.1 Data Preprocessing & Aggregation")
    add_body_paragraph(
        "Hourly timestamp records were aggregated into daily total water consumption per building. Date continuity was verified across all buildings "
        "to ensure uniform temporal spacing."
    )
    add_section_heading("3.2 Chronological Feature Engineering")
    add_body_paragraph(
        "To prevent data leakage, lag features were constructed using shift operations (shift(1) and shift(7)) within each household group. "
        "Rolling statistics (7-day mean and standard deviation) were calculated on past shifted data. Calendar features (day of week, month, "
        "day of year, and weekend indicator) were extracted from the calendar date."
    )
    add_section_heading("3.3 Chronological Dataset Split")
    add_body_paragraph(
        "In time-series forecasting, random splitting causes future data leakage. A strict chronological split was applied: 70% for training (Jan–Sep), "
        "15% for validation (Sep–Nov), and 15% for testing (Nov–Dec)."
    )
    add_section_heading("3.4 Decision Tree Regression Algorithm")
    add_body_paragraph(
        "The assigned algorithm, Decision Tree Regressor (sklearn.tree.DecisionTreeRegressor), recursively partitions the feature space to minimize "
        "Mean Squared Error (MSE) within leaf nodes. Optimal hyperparameters determined via validation tuning were max_depth=5, min_samples_split=5, "
        "and min_samples_leaf=5."
    )

    # --- CHAPTER 4 ---
    add_chapter_heading("Chapter 4: Implementation")
    add_section_heading("4.1 Environment and Libraries")
    add_body_paragraph(
        "The implementation was executed in Python 3.15 utilizing pandas, numpy, scikit-learn, matplotlib, seaborn, and streamlit."
    )
    add_section_heading("4.2 Project Architecture")
    add_body_paragraph(
        "The project follows a modular structure: data_processing.py (cleaning and feature engineering), train_model.py (model training and tuning), "
        "evaluate_model.py (metrics and evaluation plots), and app.py (Streamlit interactive interface)."
    )
    add_section_heading("4.3 Core Implementation Code Snippet")
    add_body_paragraph(
        "from sklearn.tree import DecisionTreeRegressor\n"
        "model = DecisionTreeRegressor(max_depth=5, min_samples_split=5, min_samples_leaf=5, random_state=42)\n"
        "model.fit(X_train_val, y_train_val)\n"
        "predictions = model.predict(X_test)"
    )

    # --- CHAPTER 5 ---
    add_chapter_heading("Chapter 5: Results and Discussion")
    add_section_heading("5.1 Model Performance Evaluation")
    add_body_paragraph(
        "The models were evaluated on the untouched test set (540 samples). The table below summarizes the test performance metrics."
    )

    # Evaluation table
    table_eval = doc.add_table(rows=5, cols=4)
    table_eval.alignment = WD_TABLE_ALIGNMENT.CENTER
    eval_headers = ["Model", "MAE (L)", "RMSE (L)", "R² Score"]
    for j, h in enumerate(eval_headers):
        table_eval.rows[0].cells[j].text = h
        table_eval.rows[0].cells[j].paragraphs[0].runs[0].font.bold = True

    eval_rows = [
        ("Baseline (Lag-1)", "104.46", "143.33", "-0.0144"),
        ("Baseline (Lag-7)", "114.90", "161.66", "-0.2903"),
        ("Linear Regression", "87.58", "119.42", "0.2958"),
        ("Decision Tree Regressor (Assigned)", "88.60", "121.71", "0.2686")
    ]
    for i, row_vals in enumerate(eval_rows):
        for j, val in enumerate(row_vals):
            table_eval.rows[i+1].cells[j].text = val

    for row in table_eval.rows:
        for cell in row.cells:
            set_cell_margins(cell, 80, 80, 100, 100)
            for p in cell.paragraphs:
                if p.runs:
                    p.runs[0].font.name = 'Times New Roman'
                    p.runs[0].font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    add_section_heading("5.2 Discussion of Results")
    add_body_paragraph(
        "The Decision Tree Regressor substantially outperforms naive persistence baselines (Lag-1 and Lag-7), proving that temporal lag features "
        "and rolling statistics capture meaningful consumption dynamics. Linear Regression achieves comparable metrics (MAE 87.58 vs 88.60), "
        "as daily water consumption exhibits largely linear relationships with recent lagged usage. However, the Decision Tree provides superior "
        "interpretability via explicit decision rules and feature importance rankings."
    )

    add_section_heading("5.3 Visualisations")
    add_body_paragraph("Figure 5.1: Actual vs Predicted Water Consumption")
    if os.path.exists('reports/actual_vs_predicted.png'):
        doc.add_picture('reports/actual_vs_predicted.png', width=Inches(5.5))

    add_body_paragraph("Figure 5.2: Feature Importance Chart")
    if os.path.exists('reports/feature_importance.png'):
        doc.add_picture('reports/feature_importance.png', width=Inches(5.5))

    # --- CHAPTER 6 ---
    add_chapter_heading("Chapter 6: Conclusion and Future Enhancements")
    add_section_heading("6.1 Conclusion")
    add_body_paragraph(
        "This case study successfully implemented an end-to-end data mining pipeline for household water consumption forecasting using Decision Tree "
        "Regression. Rigorous preprocessing, leak-free lag feature engineering, and validation-set tuning resulted in a robust model that outperforms "
        "persistence baselines and offers transparent decision-making."
    )
    add_section_heading("6.2 Future Enhancements")
    add_body_paragraph(
        "Future work includes integrating IoT smart-meter streaming APIs, evaluating ensemble methods such as Random Forests and Gradient Boosted Trees, "
        "and incorporating localized weather forecasts."
    )

    # --- REFERENCES ---
    add_chapter_heading("References")
    refs = [
        "1. Zenodo Dataset Record 15180141: Residential Household Dataset: Occupancy, Water, and Electricity Data. DOI: 10.5281/zenodo.15180140.",
        "2. Pedregosa, F. et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825-2830.",
        "3. McKinney, W. (2010). Data Structures for Statistical Computing in Python. Proceedings of the 9th Python in Science Conference, 56-61.",
        "4. Streamlit Documentation. Available at: https://docs.streamlit.io",
        "5. Han, J., Kamber, M., & Pei, J. (2011). Data Mining: Concepts and Techniques (3rd ed.). Morgan Kaufmann."
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(ref)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)

    doc.save('reports/Water_Consumption_Forecasting_Report.docx')
    print("Successfully generated reports/Water_Consumption_Forecasting_Report.docx")

if __name__ == '__main__':
    create_report()
