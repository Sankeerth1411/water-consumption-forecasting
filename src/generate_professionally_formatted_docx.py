import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_footer(section):
    # Header
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_before = Pt(0)
    hp.paragraph_format.space_after = Pt(0)
    hrun = hp.add_run("HOUSEHOLD WATER CONSUMPTION FORECASTING USING DECISION TREE REGRESSION")
    hrun.font.name = 'Times New Roman'
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = RGBColor(128, 128, 128)

    # Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    frun = fp.add_run("Department of Artificial Intelligence and Data Science, NMIT")
    frun.font.name = 'Times New Roman'
    frun.font.size = Pt(8.5)
    frun.font.color.rgb = RGBColor(128, 128, 128)

def create_professionally_formatted_report():
    os.makedirs('reports', exist_ok=True)
    doc = Document()

    # Page Margins (1 inch)
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    add_header_footer(section)

    # Base style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)

    def add_chapter_heading(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(24)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 51, 102) # Academic navy
        return p

    def add_section_heading(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(51, 51, 51)
        return p

    def add_subsection_heading(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(60, 60, 60)
        return p

    def add_body_paragraph(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        return p

    def add_bullet_point(text):
        p = doc.add_paragraph(style='List Bullet')
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        return p

    def add_numbered_point(number, text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.left_indent = Inches(0.25)
        run_num = p.add_run(f"{number} ")
        run_num.font.name = 'Times New Roman'
        run_num.font.size = Pt(12)
        run_num.font.bold = True
        run_text = p.add_run(text)
        run_text.font.name = 'Times New Roman'
        run_text.font.size = Pt(12)
        return p

    def add_code_block(code_text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Inches(0.4)
        run = p.add_run(code_text)
        run.font.name = 'Courier New'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(30, 30, 30)
        return p

    def add_caption(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10.5)
        run.font.italic = True
        run.font.color.rgb = RGBColor(50, 50, 50)
        return p

    # ==================== 1. COVER PAGE ====================
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(24)
    title_p.paragraph_format.space_after = Pt(12)
    r = title_p.add_run("NITTE MEENAKSHI INSTITUTE OF TECHNOLOGY\n(An Autonomous Institution, Affiliated to VTU, Belagavi)\nDepartment of Artificial Intelligence and Data Science\nYelahanka, Bengaluru - 560064")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True

    doc.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.CENTER

    proj_title_p = doc.add_paragraph()
    proj_title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    proj_title_p.paragraph_format.space_before = Pt(36)
    proj_title_p.paragraph_format.space_after = Pt(36)
    r2 = proj_title_p.add_run("A DATA MINING INDIVIDUAL PRACTICAL CASE STUDY REPORT ON\n\n\"HOUSEHOLD WATER CONSUMPTION FORECASTING USING DECISION TREE REGRESSION\"")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(18)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0, 51, 102)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_before = Pt(12)
    sub_p.paragraph_format.space_after = Pt(36)
    r3 = sub_p.add_run("Submitted in partial fulfillment of the requirements for the course\nDATA MINING (Course Code: DATA MINING-22ADG53)\n\nAcademic Year: 2026–2027")
    r3.font.name = 'Times New Roman'
    r3.font.size = Pt(12)

    # Details table
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    details = [
        ("Submitted By:", "R. Sankeerth (USN: 1NT24AD047)"),
        ("Semester / Section:", "5th Semester / Artificial Intelligence & Data Science"),
        ("Assigned Algorithm:", "Decision Tree Regressor (sklearn.tree.DecisionTreeRegressor)"),
        ("Faculty Guide:", "Mrs Nethravathi T L, Assistant Professor G-1"),
        ("Department:", "Department of Artificial Intelligence and Data Science")
    ]
    for i, (k, v) in enumerate(details):
        row = table.rows[i]
        row.cells[0].text = k
        row.cells[1].text = v
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        for cell in row.cells:
            set_cell_margins(cell, 80, 80, 100, 100)
            for p in cell.paragraphs:
                if p.runs:
                    p.runs[0].font.name = 'Times New Roman'
                    p.runs[0].font.size = Pt(11)

    doc.add_page_break()

    # ==================== 2. INSTITUTIONAL FRONT MATTER ====================
    add_chapter_heading("Institutional Vision, Mission, PEOs & PSOs")

    add_section_heading("Nitte Meenakshi Institute of Technology (NMIT)")
    add_subsection_heading("Vision of the Institute")
    add_body_paragraph(
        "To provide quality technical education through research, innovation, and teamwork for the development of society."
    )
    add_subsection_heading("Mission of the Institute")
    add_bullet_p1 = add_bullet_point("M1: To impart quality technical education by nurturing a conducive learning environment.")
    add_bullet_p2 = add_bullet_point("M2: To foster research and innovation through collaborations with industries and academic institutions.")
    add_bullet_p3 = add_bullet_point("M3: To instill ethical values, leadership qualities, and social responsibilities in students.")

    add_section_heading("Department of Artificial Intelligence and Data Science")
    add_subsection_heading("Vision of the Department")
    add_body_paragraph(
        "To be a center of excellence in Artificial Intelligence and Data Science education and research, empowering students to solve real-world challenges."
    )
    add_subsection_heading("Mission of the Department")
    add_bullet_point("DM1: To provide robust theoretical and practical foundations in AI and Data Science.")
    add_bullet_point("DM2: To encourage multidisciplinary research, innovation, and industry partnerships.")
    add_bullet_point("DM3: To produce socially responsible professionals equipped with ethical principles.")

    add_section_heading("Programme Educational Objectives (PEOs)")
    add_bullet_point("PEO1: Graduates will establish themselves as proficient professionals in Artificial Intelligence, Data Science, and allied domains.")
    add_bullet_point("PEO2: Graduates will engage in lifelong learning, advanced higher studies, and cutting-edge research.")
    add_bullet_point("PEO3: Graduates will exhibit strong leadership, communication, teamwork, and professional ethics.")

    add_section_heading("Programme Specific Outcomes (PSOs)")
    add_bullet_point("PSO1: Apply core concepts of machine learning, data mining, and statistical modeling to analyze complex datasets.")
    add_bullet_point("PSO2: Design, develop, and deploy intelligent, scalable AI and data-driven software solutions for real-world applications.")

    doc.add_page_break()

    # ==================== 3. TABLE OF CONTENTS ====================
    add_chapter_heading("Table of Contents")
    toc_items = [
        "Abstract .................................................................................................................... i",
        "Chapter 1: Introduction .............................................................................................. 1",
        "Chapter 2: Dataset Description .................................................................................. 4",
        "Chapter 3: Methodology and Algorithm ................................................................... 7",
        "Chapter 4: Implementation .................................................................................... 12",
        "Chapter 5: Results and Discussion .......................................................................... 16",
        "Chapter 6: Conclusion and Future Enhancements ................................................. 21",
        "References ............................................................................................................... 23"
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)

    doc.add_page_break()

    # ==================== 4. ABSTRACT ====================
    add_chapter_heading("Abstract")
    add_body_paragraph(
        "Accurate household water consumption forecasting is vital for municipal water resource management, demand response planning, "
        "and sustainable urban infrastructure operations. This individual practical case study formulates daily water consumption estimation "
        "as a supervised machine learning regression task using historical residential monitoring data. Following rigorous data cleaning, temporal "
        "aggregation, and leakage-free feature engineering—incorporating lagged consumption and rolling window statistics—a Decision Tree "
        "Regressor (sklearn.tree.DecisionTreeRegressor) was trained, hyperparameter-tuned on a validation set, and evaluated against naive "
        "persistence baselines and linear regression on an untouched chronological test set. The model achieved a Test Mean Absolute Error (MAE) "
        "of 88.60 Liters and Root Mean Squared Error (RMSE) of 121.71 Liters, successfully outperforming historical persistence baselines while providing "
        "transparent, interpretable tree decision rules. To demonstrate practical utility, a lightweight Streamlit interactive web application "
        "was developed, featuring real-time forecasting and leaf-node diagnostic introspection. This report documents the complete end-to-end "
        "data mining lifecycle, experimental results, and domain insights."
    )

    doc.add_page_break()

    # ==================== CHAPTER 1 ====================
    add_chapter_heading("Chapter 1: Introduction")

    add_section_heading("1.1 Background and Introduction")
    add_body_paragraph(
        "Modern urban centers face unprecedented challenges in managing municipal water resources efficiently. Rapid urbanization, population "
        "growth, climate change, and dwindling groundwater reserves place immense pressure on municipal authorities to optimize water distribution "
        "networks. Traditional water management systems often operate reactively, leading to supply-demand imbalances, water wastage, and delayed "
        "leak detection. Intelligent water consumption forecasting provides a proactive solution by leveraging historical usage patterns, temporal "
        "cycles, and household occupancy dynamics to predict future demand accurately."
    )
    add_body_paragraph(
        "In the realm of data mining and predictive analytics, time-series forecasting of utility consumption allows municipalities and smart "
        "water utilities to optimize pumping schedules, maintain balanced reservoir storage, and implement dynamic pricing structures. Furthermore, "
        "accurate consumption forecasts empower individual households to monitor their usage, detect anomalous leaks, and achieve sustainable "
        "domestic water conservation."
    )

    add_section_heading("1.2 Real-World Importance of Water Consumption Forecasting")
    add_body_paragraph(
        "The practical significance of forecasting domestic water consumption spans multiple operational domains:"
    )
    add_numbered_point("1.", "Municipal Infrastructure Optimization: Water treatment plants and distribution pumping stations consume vast amounts of electrical energy. Accurate day-ahead forecasts enable operators to schedule pumping during off-peak electricity tariff hours, substantially reducing municipal operational costs.")
    add_numbered_point("2.", "Leak and Anomaly Detection: Significant deviations between forecasted consumption and actual metered readings serve as early indicators of pipe bursts, hidden leaks, or faulty water meters, preventing structural damage and water loss.")
    add_numbered_point("3.", "Demand-Side Management: Empowering consumers with localized forecasts and consumption insights encourages behavioral changes toward sustainable water usage.")

    add_section_heading("1.3 Problem Statement")
    add_body_paragraph(
        "Given historical daily water consumption time-series data and occupancy indicator features across residential households, the objective "
        "is to construct a robust supervised machine learning regression model using Decision Tree Regression to predict the next calendar day's "
        "total water consumption (in liters) for a given household, utilizing only information available prior to the prediction date without "
        "introducing data leakage."
    )

    add_section_heading("1.4 Objectives of the Case Study")
    add_body_paragraph(
        "The primary objectives of this individual practical case study are:"
    )
    add_bullet_point("To acquire, inspect, and audit the public Zenodo residential water consumption and occupancy dataset (Record 15180141).")
    add_bullet_point("To design and execute a reproducible data preprocessing and aggregation pipeline, converting high-resolution timestamps into daily consumption totals.")
    add_bullet_point("To perform leak-free chronological feature engineering, constructing lagged consumption variables and rolling window statistics.")
    add_bullet_point("To train, hyperparameter-tune, and evaluate a Decision Tree Regressor (sklearn.tree.DecisionTreeRegressor) adhering strictly to chronological train, validation, and test splits.")
    add_bullet_point("To benchmark the Decision Tree against naive persistence baselines and linear regression.")
    add_bullet_point("To develop an interactive Streamlit demonstration application for real-time forecasting and model introspection.")

    add_section_heading("1.5 Scope of the Study")
    add_body_paragraph(
        "This study focuses on 10 residential households monitored throughout the year 2022. The analytical scope covers data cleaning, temporal "
        "aggregation, feature construction, model training, hyperparameter optimization, test evaluation, and web application deployment. External "
        "meteorological data was evaluated but excluded due to spatial anonymization constraints; instead, robust temporal lags and seasonal "
        "indicators were utilized."
    )

    add_section_heading("1.6 Practical Applications")
    add_body_paragraph(
        "The methodologies developed in this case study are directly applicable to smart city water grids, IoT smart-metering platforms, domestic "
        "conservation mobile applications, and municipal utility resource planning."
    )

    doc.add_page_break()

    # ==================== CHAPTER 2 ====================
    add_chapter_heading("Chapter 2: Dataset Description")

    add_section_heading("2.1 Dataset Introduction")
    add_body_paragraph(
        "Data selection is a critical foundational step in any data mining project. For this case study, a rigorous vetting process was conducted "
        "to identify a dataset containing authentic residential water consumption records, occupancy indicators, and sufficient temporal depth "
        "to support robust time-series forecasting."
    )

    add_section_heading("2.2 Original Dataset Source and Citation")
    add_body_paragraph(
        "The dataset utilized is the 'Residential Household Dataset: Occupancy, Water, and Electricity Data' hosted publicly on Zenodo. "
        "The original data source is formally cited as:"
    )
    add_bullet_point("Repository: Zenodo")
    add_bullet_point("Record ID: 15180141")
    add_bullet_point("DOI: 10.5281/zenodo.15180140")
    add_bullet_point("License: Creative Commons Attribution (CC-BY)")

    add_section_heading("2.3 Dataset Characteristics")
    add_body_paragraph(
        "The curated dataset contains records spanning from January 1, 2022, to December 30, 2022. It covers 10 distinct residential buildings "
        "(BUILD-D-01 through BUILD-D-10). The raw release contains hourly timestamped observations, which were aggregated into daily total "
        "consumption records comprising 3,640 clean daily rows."
    )

    add_section_heading("2.4 Description of Attributes (Data Dictionary)")
    add_body_paragraph("Table 2.1 outlines the schema and data dictionary of the processed dataset.")

    # Data Dictionary Table
    table_dict = doc.add_table(rows=10, cols=4)
    table_dict.alignment = WD_TABLE_ALIGNMENT.CENTER
    dict_headers = ["Attribute Name", "Data Type", "Unit", "Description"]
    for j, h in enumerate(dict_headers):
        table_dict.rows[0].cells[j].text = h
        table_dict.rows[0].cells[j].paragraphs[0].runs[0].font.bold = True

    dict_rows = [
        ("building", "String (Categorical)", "NID", "Unique building identifier (BUILD-D-01 to BUILD-D-10)"),
        ("date", "Date (YYYY-MM-DD)", "Date", "Calendar date of observation"),
        ("water_use_l", "Float", "Liters (L)", "Total water consumption on target day"),
        ("lag_1", "Float", "Liters (L)", "Water consumption 1 day prior (t-1) [Lag feature]"),
        ("lag_7", "Float", "Liters (L)", "Water consumption 7 days prior (t-7) [Lag feature]"),
        ("rolling_mean_7", "Float", "Liters (L)", "Rolling 7-day average water consumption"),
        ("rolling_std_7", "Float", "Liters (L)", "Rolling 7-day standard deviation of consumption"),
        ("day_of_week", "Integer", "0-6", "Day of the week (0 = Monday, 6 = Sunday)"),
        ("occ_01 to occ_05", "Float", "Index (0-1)", "Household occupancy indicators")
    ]
    for i, row_vals in enumerate(dict_rows):
        for j, val in enumerate(row_vals):
            table_dict.rows[i+1].cells[j].text = val

    for row in table_dict.rows:
        for cell in row.cells:
            set_cell_margins(cell, 80, 80, 100, 100)
            for p in cell.paragraphs:
                if p.runs:
                    p.runs[0].font.name = 'Times New Roman'
                    p.runs[0].font.size = Pt(10)

    add_caption("Table 2.1: Data Dictionary of Processed Residential Dataset")

    add_section_heading("2.5 Target Variable")
    add_body_paragraph(
        "The primary target variable is water_use_l, representing the total daily water consumption expressed in Liters (L). For forecasting, "
        "the model predicts the target value for calendar day t using features constructed exclusively from days t-1 and earlier."
    )

    add_section_heading("2.6 Sample Dataset Records")
    add_body_paragraph("Table 2.2 presents actual sample records extracted from the processed dataset.")

    # Sample Records Table
    table_samp = doc.add_table(rows=4, cols=5)
    table_samp.alignment = WD_TABLE_ALIGNMENT.CENTER
    samp_headers = ["Building", "Date", "Target Water (L)", "Lag-1 Water (L)", "Rolling Mean (L)"]
    for j, h in enumerate(samp_headers):
        table_samp.rows[0].cells[j].text = h
        table_samp.rows[0].cells[j].paragraphs[0].runs[0].font.bold = True

    samp_rows = [
        ("BUILD-D-01", "2022-01-08", "275.00", "312.40", "295.10"),
        ("BUILD-D-01", "2022-01-09", "296.00", "275.00", "290.30"),
        ("BUILD-D-01", "2022-01-10", "240.00", "296.00", "288.70")
    ]
    for i, row_vals in enumerate(samp_rows):
        for j, val in enumerate(row_vals):
            table_samp.rows[i+1].cells[j].text = val

    for row in table_samp.rows:
        for cell in row.cells:
            set_cell_margins(cell, 80, 80, 100, 100)
            for p in cell.paragraphs:
                if p.runs:
                    p.runs[0].font.name = 'Times New Roman'
                    p.runs[0].font.size = Pt(10)

    add_caption("Table 2.2: Sample Records from Processed Dataset")

    add_section_heading("2.7 Reason for Dataset Selection")
    add_body_paragraph(
        "This dataset was chosen because it provides authentic, multi-household metered consumption data with clean temporal indexing and zero "
        "missing values in its curated release, making it ideal for academic benchmarking and machine learning experimentation."
    )

    doc.add_page_break()

    # ==================== CHAPTER 3 ====================
    add_chapter_heading("Chapter 3: Methodology and Algorithm")

    add_section_heading("3.1 Proposed Methodology")
    add_body_paragraph(
        "The end-to-end data mining methodology follows a structured pipeline: Data Ingestion and Aggregation, Leak-Free Feature Engineering, "
        "Chronological Dataset Splitting, Model Training with Hyperparameter Optimization, Performance Evaluation against Baselines, and "
        "Streamlit Application Deployment."
    )

    add_section_heading("3.2 Data Preprocessing & Cleaning")
    add_body_paragraph(
        "Raw hourly records were ingested and verified for timestamp continuity. Occupancy variables were converted to numeric types. "
        "Hourly water use values were summed per building per calendar day to form daily aggregate records."
    )

    add_section_heading("3.3 Feature Engineering and Lag Construction")
    add_body_paragraph(
        "To capture temporal autocorrelation, lag features were generated by shifting consumption series within each building group: "
        "lag_1 (previous day consumption) and lag_7 (consumption one week prior). Rolling statistics (7-day mean and standard deviation) "
        "were computed on past shifted data to capture rolling trends."
    )

    add_section_heading("3.4 Chronological Dataset Splitting")
    add_body_paragraph(
        "To prevent look-ahead bias and future data leakage, a strict chronological split was enforced: 70% training (Jan 8 to Sep 13), "
        "15% validation (Sep 14 to Nov 6), and 15% testing (Nov 7 to Dec 30)."
    )

    add_section_heading("3.5 Introduction to Decision Tree Regression")
    add_body_paragraph(
        "Decision Tree Regression is a non-parametric supervised learning method that predicts numerical targets by learning simple decision rules "
        "inferred from data features. The assigned algorithm, sklearn.tree.DecisionTreeRegressor, recursively splits data to minimize node variance."
    )

    add_section_heading("3.6 Mathematical Explanation of Decision Tree Splitting")
    add_body_paragraph(
        "At each node m representing region Rm with Nm observations, the impurity is measured by Mean Squared Error (MSE):\n"
        "H(Rm) = (1 / Nm) * sum_{i in Rm} (y_i - c_m)^2\n"
        "where c_m is the mean target value of the training observations in region Rm:\n"
        "c_m = (1 / Nm) * sum_{i in Rm} y_i\n"
        "The algorithm selects a feature j and threshold s to split region Rm into R_left(j, s) and R_right(j, s) minimizing total child impurity."
    )

    add_section_heading("3.7 Hyperparameter Selection")
    add_body_paragraph(
        "Hyperparameters were tuned on the validation set. Optimal parameters identified were: max_depth = 5, min_samples_split = 5, and "
        "min_samples_leaf = 5."
    )

    add_section_heading("3.8 Evaluation Metrics")
    add_body_paragraph(
        "Model performance was evaluated using Mean Absolute Error (MAE), Root Mean Squared Error (RMSE), and Coefficient of Determination (R²)."
    )

    doc.add_page_break()

    # ==================== CHAPTER 4 ====================
    add_chapter_heading("Chapter 4: Implementation")

    add_section_heading("4.1 Hardware and Software Requirements")
    add_body_paragraph(
        "Hardware: Standard student laptop with CPU (Intel Core / AMD Ryzen, 8GB+ RAM).\n"
        "Software: Windows 11, Python 3.15, Visual Studio Code."
    )

    add_section_heading("4.2 Libraries and Frameworks")
    add_body_paragraph(
        "Core libraries utilized include pandas (data manipulation), numpy (numerical arrays), scikit-learn (Decision Tree and metrics), "
        "matplotlib and seaborn (visualizations), and streamlit (web application)."
    )

    add_section_heading("4.3 Project Architecture & Directory Structure")
    add_body_paragraph(
        "The project is organized into raw/processed data folders, src/ modules (data_processing.py, train_model.py, evaluate_model.py), "
        "models/ (saved joblib artifact), reports/ (evaluation CSV and charts), and app.py (Streamlit application)."
    )

    add_section_heading("4.4 Core Implementation Code Snippets")
    add_body_paragraph("Snippet 4.1: Model Training (`src/train_model.py`)")
    add_code_block(
        "from sklearn.tree import DecisionTreeRegressor\n"
        "import joblib\n\n"
        "# Initialize Decision Tree Regressor with optimal hyperparameters\n"
        "dt_model = DecisionTreeRegressor(\n"
        "    max_depth=5,\n"
        "    min_samples_split=5,\n"
        "    min_samples_leaf=5,\n"
        "    random_state=42\n"
        ")\n"
        "dt_model.fit(X_train_val, y_train_val)\n"
        "joblib.dump({'model': dt_model, 'feature_cols': feature_cols}, 'models/decision_tree_model.joblib')"
    )

    add_body_paragraph("Snippet 4.2: Streamlit Application Prediction (`app.py`)")
    add_code_block(
        "import streamlit as st\n"
        "import joblib\n"
        "import pandas as pd\n\n"
        "artifact = joblib.load('models/decision_tree_model.joblib')\n"
        "model = artifact['model']\n"
        "feature_cols = artifact['feature_cols']\n\n"
        "input_df = pd.DataFrame([user_inputs])[feature_cols]\n"
        "prediction = model.predict(input_df)[0]\n"
        "leaf_id = model.apply(input_df)[0]"
    )

    doc.add_page_break()

    # ==================== CHAPTER 5 ====================
    add_chapter_heading("Chapter 5: Results and Discussion")

    add_section_heading("5.1 Model Performance Comparison")
    add_body_paragraph(
        "The models were evaluated on the untouched chronological test set (540 samples). Table 5.1 compares the Decision Tree Regressor "
        "against persistence baselines and Linear Regression."
    )

    # Evaluation Table
    table_res = doc.add_table(rows=5, cols=4)
    table_res.alignment = WD_TABLE_ALIGNMENT.CENTER
    res_headers = ["Model", "MAE (L)", "RMSE (L)", "R² Score"]
    for j, h in enumerate(res_headers):
        table_res.rows[0].cells[j].text = h
        table_res.rows[0].cells[j].paragraphs[0].runs[0].font.bold = True

    res_rows = [
        ("Baseline (Lag-1)", "104.46", "143.33", "-0.0144"),
        ("Baseline (Lag-7)", "114.90", "161.66", "-0.2903"),
        ("Linear Regression", "87.58", "119.42", "0.2958"),
        ("Decision Tree Regressor (Assigned)", "88.60", "121.71", "0.2686")
    ]
    for i, row_vals in enumerate(res_rows):
        for j, val in enumerate(row_vals):
            table_res.rows[i+1].cells[j].text = val

    for row in table_res.rows:
        for cell in row.cells:
            set_cell_margins(cell, 80, 80, 100, 100)
            for p in cell.paragraphs:
                if p.runs:
                    p.runs[0].font.name = 'Times New Roman'
                    p.runs[0].font.size = Pt(10)

    add_caption("Table 5.1: Performance Comparison Across Test Set Models")

    add_section_heading("5.2 Discussion of Findings & Algorithm Comparison")
    add_body_paragraph(
        "The experimental results demonstrate that the Decision Tree Regressor successfully outperforms naive persistence baselines (Lag-1 and Lag-7), "
        "confirming the predictive value of engineered lag and rolling features. Linear Regression achieves a slightly lower MAE (87.58 L vs 88.60 L) "
        "and higher R² (0.2958 vs 0.2686). This occurs because daily water consumption exhibits predominantly linear relationships with immediate "
        "past usage. However, while Linear Regression provides coefficients, the Decision Tree Regressor provides superior explainability through "
        "hierarchical decision rules and feature importance rankings."
    )

    add_section_heading("5.3 Visualisations")
    add_body_paragraph("Figure 5.1 illustrates the actual versus predicted daily water consumption across the test set.")
    if os.path.exists('reports/actual_vs_predicted.png'):
        doc.add_picture('reports/actual_vs_predicted.png', width=Inches(5.5))
    add_caption("Figure 5.1: Actual vs Predicted Daily Water Consumption")

    add_body_paragraph("Figure 5.2 displays the feature importance ranking derived from the trained Decision Tree model.")
    if os.path.exists('reports/feature_importance.png'):
        doc.add_picture('reports/feature_importance.png', width=Inches(5.5))
    add_caption("Figure 5.2: Decision Tree Feature Importance Ranking")

    doc.add_page_break()

    # ==================== CHAPTER 6 ====================
    add_chapter_heading("Chapter 6: Conclusion and Future Enhancements")

    add_section_heading("6.1 Summary of Work")
    add_body_paragraph(
        "This individual practical case study successfully executed an end-to-end data mining project for household water consumption forecasting "
        "using Decision Tree Regression. Rigorous data cleaning, leak-free lag feature engineering, and validation-set hyperparameter tuning "
        "established a robust predictive pipeline."
    )

    add_section_heading("6.2 Major Findings")
    add_body_paragraph(
        "The Decision Tree model achieved a Test MAE of 88.60 Liters, outperforming historical persistence baselines and confirming that recent "
        "consumption history is the strongest predictor of future water demand."
    )

    add_section_heading("6.3 Limitations")
    add_body_paragraph(
        "Limitations include the anonymized nature of the residential dataset preventing external weather API enrichment, and piecewise-constant "
        "quantization inherent to decision trees."
    )

    add_section_heading("6.4 Future Enhancements")
    add_body_paragraph(
        "Future work will explore ensemble methods such as Random Forests and Gradient Boosted Trees, alongside IoT smart-meter streaming integrations."
    )

    doc.add_page_break()

    # ==================== REFERENCES ====================
    add_chapter_heading("References")
    refs = [
        "1. Zenodo Dataset Record 15180141: Residential Household Dataset: Occupancy, Water, and Electricity Data. DOI: 10.5281/zenodo.15180140.",
        "2. Pedregosa, F. et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825-2830.",
        "3. McKinney, W. (2010). Data Structures for Statistical Computing in Python. Proceedings of the 9th Python in Science Conference, 56-61.",
        "4. Hunter, J. D. (2007). Matplotlib: A 2D Graphics Environment. Computing in Science & Engineering, 9(3), 90-95.",
        "5. Streamlit Documentation. Available at: https://docs.streamlit.io",
        "6. Han, J., Kamber, M., & Pei, J. (2011). Data Mining: Concepts and Techniques (3rd ed.). Morgan Kaufmann."
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(ref)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)

    output_path = 'reports/Water_Consumption_Forecasting_Report_Professionally_Formatted.docx'
    doc.save(output_path)
    print(f"Successfully generated professionally formatted report at {output_path}")

if __name__ == '__main__':
    create_professionally_formatted_report()
